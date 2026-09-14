from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from src.doc_generator import DocGenerator
from src.models import CandidateProfile, Experience, OptimizedContent


def load_profile() -> CandidateProfile:
    return CandidateProfile.from_yaml(Path(__file__).parent.parent / "data" / "profile.yaml")


def make_optimized(profile: CandidateProfile) -> OptimizedContent:
    return OptimizedContent(
        matched_keywords=["VLAN", "DHCP", "Cisco"],
        match_score=80,
        prioritized_competence_categories=["routage_commutation", "securite_reseau"],
        rewritten_summary="Résumé adapté à l'offre, sans invention.",
        rewritten_experiences=[
            Experience(
                poste="Projet BTS — Déploiement d'un serveur DHCP",
                entreprise="Projet d'études",
                periode="2025",
                lieu="",
                bullets=["Configuration DHCP multi-VLAN sur routeurs Cisco"],
            )
        ],
    )


def test_generated_docx_has_single_column_and_no_tables(tmp_path: Path) -> None:
    profile = load_profile()
    optimized = make_optimized(profile)
    generator = DocGenerator(output_dir=tmp_path)

    path = generator.build(profile, optimized)

    assert path.exists()
    document = Document(path)

    # ATS : aucun tableau
    assert document.tables == []

    # ATS : aucune image / zone de texte
    assert len(document.inline_shapes) == 0

    # ATS : une seule colonne (pas de balise w:cols avec num > 1)
    section = document.sections[0]
    cols_el = section._sectPr.find(qn("w:cols"))
    if cols_el is not None:
        assert cols_el.get(qn("w:num")) in (None, "1")

    # Police standard cohérente
    assert document.styles["Normal"].font.name == "Calibri"


def test_generated_docx_text_contains_matched_keywords(tmp_path: Path) -> None:
    profile = load_profile()
    optimized = make_optimized(profile)
    generator = DocGenerator(output_dir=tmp_path)

    path = generator.build(profile, optimized)
    document = Document(path)
    full_text = "\n".join(p.text for p in document.paragraphs)

    assert "DHCP" in full_text
    assert profile.full_name in full_text
    assert optimized.rewritten_summary in full_text


def test_build_without_optimization_falls_back_to_base_profile(tmp_path: Path) -> None:
    profile = load_profile()
    generator = DocGenerator(output_dir=tmp_path)

    path = generator.build(profile)
    document = Document(path)
    full_text = "\n".join(p.text for p in document.paragraphs)

    assert profile.summary.strip() in full_text
