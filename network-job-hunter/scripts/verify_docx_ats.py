#!/usr/bin/env python3
"""Script de vérification manuelle (phase VERIFY) :

Génère un CV à partir d'une offre fictive "Technicien Réseau en alternance"
(sans appeler l'API Anthropic — un OptimizedContent factice est utilisé),
puis imprime un rapport de conformité ATS : nombre de tableaux, nombre de
colonnes, police utilisée, absence d'images, présence des mots-clés,
et tente la conversion en PDF.
"""
from __future__ import annotations

import sys
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from config import settings  # noqa: E402
from src.doc_generator import DocGenerator  # noqa: E402
from src.models import CandidateProfile, Experience, JobOffer, OptimizedContent  # noqa: E402


def fake_job() -> JobOffer:
    return JobOffer(
        job_id="verify-1",
        title="Technicien Réseau en alternance",
        company="NetSecure Solutions",
        description=(
            "Configuration de VLAN, gestion de serveurs DHCP, routage RIP/OSPF, "
            "supervision Zabbix, premières briques de sécurité réseau (pare-feu, WAF)."
        ),
        url="https://example.com/offres/verify-1",
        date_posted=date.today() - timedelta(days=2),
        location="Paris (75)",
        source="verify_script",
    )


def fake_optimized(profile: CandidateProfile) -> OptimizedContent:
    return OptimizedContent(
        matched_keywords=["VLAN", "DHCP", "RIP", "Zabbix", "Cisco"],
        match_score=84,
        prioritized_competence_categories=["routage_commutation", "supervision_exploitation"],
        rewritten_summary=(
            "Étudiant en BTS SIO SISR avec une expérience concrète en configuration "
            "VLAN/DHCP sur routeurs Cisco et en supervision réseau."
        ),
        rewritten_experiences=profile.experiences,
    )


def print_report(docx_path: Path) -> bool:
    document = Document(docx_path)
    ok = True

    def check(label: str, condition: bool) -> None:
        nonlocal ok
        status = "OK " if condition else "FAIL"
        print(f"  [{status}] {label}")
        ok = ok and condition

    print(f"\nRapport de conformité ATS — {docx_path.name}")
    check("Aucun tableau", document.tables == [])
    check("Aucune image / zone de texte (inline shapes)", len(document.inline_shapes) == 0)

    section = document.sections[0]
    cols_el = section._sectPr.find(qn("w:cols"))
    single_column = cols_el is None or cols_el.get(qn("w:num")) in (None, "1")
    check("Une seule colonne", single_column)

    font_name = document.styles["Normal"].font.name
    check(f"Police standard définie ({font_name})", font_name in {"Calibri", "Arial", "Times New Roman"})

    full_text = "\n".join(p.text for p in document.paragraphs)
    for keyword in ["VLAN", "DHCP", "Wesley Amavi"]:
        check(f"Mot-clé présent : '{keyword}'", keyword in full_text)

    check("Texte non vide et extractible", len(full_text.strip()) > 200)

    return ok


def main() -> None:
    profile = CandidateProfile.from_yaml(settings.PROFILE_PATH)
    job = fake_job()
    optimized = fake_optimized(profile)

    generator = DocGenerator(output_dir=settings.OUTPUT_DIR)
    docx_path = generator.build(profile, optimized, job)
    print(f"CV généré : {docx_path}")

    ok = print_report(docx_path)

    try:
        pdf_path = generator.convert_to_pdf(docx_path)
        print(f"\nConversion PDF réussie : {pdf_path} ({pdf_path.stat().st_size} octets)")
    except RuntimeError as exc:
        print(f"\n[WARN] Conversion PDF impossible : {exc}")

    print("\n" + ("Résultat global : CONFORME ATS" if ok else "Résultat global : NON CONFORME — voir FAIL ci-dessus"))
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
