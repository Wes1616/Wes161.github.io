"""Génère un CV strictement compatible ATS : une colonne, texte brut, sans
tableau, sans zone de texte, sans image, police standard. Puis convertit le
.docx en .pdf via LibreOffice en ligne de commande.
"""
from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from config import settings
from src.models import CandidateProfile, Experience, JobOffer, OptimizedContent

_FONT_NAME = "Calibri"
_BULLET = "•"

_CATEGORY_LABELS = {
    "routage_commutation": "Routage & Commutation",
    "securite_reseau": "Sécurité réseau",
    "supervision_exploitation": "Supervision & Exploitation",
    "outils": "Outils",
}


class DocGenerator:
    def __init__(self, output_dir: Path = settings.OUTPUT_DIR) -> None:
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def build(
        self,
        profile: CandidateProfile,
        optimized: OptimizedContent | None = None,
        job: JobOffer | None = None,
    ) -> Path:
        summary = optimized.rewritten_summary if optimized else profile.summary
        experiences = optimized.rewritten_experiences if optimized else profile.experiences
        category_order = (
            optimized.prioritized_competence_categories
            if optimized
            else list(profile.competences.keys())
        )

        document = Document()
        self._configure_base_style(document)

        self._add_header(document, profile)
        self._add_section_heading(document, "Profil")
        document.add_paragraph(summary)

        self._add_section_heading(document, "Compétences")
        for category in self._ordered_categories(category_order, profile.competences):
            skills = profile.competences.get(category, [])
            if not skills:
                continue
            label = _CATEGORY_LABELS.get(category, category.replace("_", " ").title())
            heading = document.add_paragraph()
            heading.add_run(f"{label} : ").bold = True
            heading.add_run(", ".join(skills))

        self._add_section_heading(document, "Expérience")
        for experience in experiences:
            self._add_experience(document, experience)

        self._add_section_heading(document, "Formation")
        for formation in profile.formation:
            para = document.add_paragraph()
            para.add_run(f"{formation.diplome} — {formation.etablissement}").bold = True
            document.add_paragraph(formation.periode)
            if formation.details:
                document.add_paragraph(formation.details)

        if profile.certifications:
            self._add_section_heading(document, "Certifications")
            for cert in profile.certifications:
                document.add_paragraph(f"{_BULLET} {cert}")

        if profile.langues:
            self._add_section_heading(document, "Langues")
            langues_line = ", ".join(f"{l['langue']} ({l['niveau']})" for l in profile.langues)
            document.add_paragraph(langues_line)

        output_path = self.output_dir / self._filename_for(profile, job)
        document.save(output_path)
        return output_path

    def convert_to_pdf(self, docx_path: Path) -> Path:
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            raise RuntimeError(
                "LibreOffice ('soffice') introuvable. Installe LibreOffice pour la "
                "conversion DOCX -> PDF, ou utilise docx2pdf sous Windows/Mac."
            )
        result = subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(docx_path.parent),
                str(docx_path),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        pdf_path = docx_path.with_suffix(".pdf")
        if not pdf_path.exists():
            raise RuntimeError(
                "La conversion PDF via LibreOffice a échoué silencieusement "
                f"(fichier attendu introuvable : {pdf_path}).\n"
                f"stdout: {result.stdout}\nstderr: {result.stderr}"
            )
        return pdf_path

    # -- internes -------------------------------------------------------

    def _configure_base_style(self, document: Document) -> None:
        style = document.styles["Normal"]
        style.font.name = _FONT_NAME
        style.font.size = Pt(11)
        section = document.sections[0]
        section.left_margin = section.right_margin = Pt(50)

    def _add_header(self, document: Document, profile: CandidateProfile) -> None:
        name_para = document.add_paragraph()
        name_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = name_para.add_run(profile.full_name)
        run.bold = True
        run.font.size = Pt(18)

        title_para = document.add_paragraph(profile.target_title)
        title_para.runs[0].font.size = Pt(12)

        contact_bits = [
            b
            for b in [profile.email, profile.phone, profile.location, profile.linkedin, profile.github]
            if b
        ]
        document.add_paragraph(" | ".join(contact_bits))

    def _add_section_heading(self, document: Document, title: str) -> None:
        heading = document.add_heading(title, level=1)
        heading.style = document.styles["Heading 1"]

    def _add_experience(self, document: Document, experience: Experience) -> None:
        title_para = document.add_paragraph()
        title_para.add_run(f"{experience.poste} — {experience.entreprise}").bold = True
        meta_bits = [b for b in [experience.periode, experience.lieu] if b]
        if meta_bits:
            document.add_paragraph(" | ".join(meta_bits))
        for bullet in experience.bullets:
            document.add_paragraph(f"{_BULLET} {bullet}")

    @staticmethod
    def _ordered_categories(priority: list[str], all_categories: dict[str, list[str]]) -> list[str]:
        ordered = [c for c in priority if c in all_categories]
        remaining = [c for c in all_categories if c not in ordered]
        return ordered + remaining

    @staticmethod
    def _filename_for(profile: CandidateProfile, job: JobOffer | None) -> str:
        base = f"CV_{profile.full_name}"
        if job:
            base += f"_{job.company}"
        slug = re.sub(r"[^A-Za-z0-9]+", "_", base).strip("_")
        return f"{slug}.docx"
